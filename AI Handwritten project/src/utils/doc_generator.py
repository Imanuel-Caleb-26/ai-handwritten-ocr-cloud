"""
Document export utilities — generates .docx and .pdf from recognized text.
"""

import os
from docx import Document
from docx.shared import Pt, RGBColor
from fpdf import FPDF


def generate_docx(text: str, output_path: str) -> str:
    """Generate a .docx file from recognized text."""
    doc = Document()

    # Title
    title = doc.add_heading("Handwriting Recognition Output", level=1)
    title.runs[0].font.color.rgb = RGBColor(0x1A, 0x73, 0xE8)

    # Body — preserve original line breaks
    for line in text.split("\n"):
        p = doc.add_paragraph(line)
        if p.runs:
            p.runs[0].font.size = Pt(12)

    doc.save(output_path)
    return output_path


class _PDF(FPDF):
    def header(self):
        self.set_font("Arial", "B", 12)
        self.set_text_color(26, 115, 232)   # Brand blue
        self.cell(0, 10, "Handwriting Recognition Output", border=0, ln=1, align="C")
        self.set_text_color(0, 0, 0)
        self.ln(4)

    def footer(self):
        self.set_y(-15)
        self.set_font("Arial", "I", 8)
        self.set_text_color(128, 128, 128)
        self.cell(0, 10, f"Page {self.page_no()}", border=0, align="C")


def generate_pdf(text: str, output_path: str) -> str:
    """Generate a .pdf file from recognized text."""
    pdf = _PDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.set_text_color(0, 0, 0)

    # Handle encoding — FPDF is Latin-1 based
    safe_text = text.encode("latin-1", errors="replace").decode("latin-1")
    pdf.multi_cell(0, 8, safe_text)

    pdf.output(output_path)
    return output_path
